from __future__ import annotations
from projectmanager.i18n import tr as _tr
from projectmanager.core.shared import *
from projectmanager.core.shared import (
    _dt
)
from projectmanager.presentation.dialogs import *

class ProjectOperationsMixin:
        def _show_new_project_wizard(self) -> None:
            """Open de v1.5 wizard voor het aanmaken van een nieuw project."""
            default_root = self.root_path_var.get().strip() or str(Path.home())
            dialog = NewProjectWizardDialog(self.root, default_root=default_root)
            self.root.wait_window(dialog.window)

            if not dialog.result:
                return

            project_path: Path = dialog.result["path"]
            add_to_locations = bool(dialog.result.get("add_to_locations", True))
            open_after = bool(dialog.result.get("open_after", False))

            try:
                info = self.analyzer.analyze(project_path, deep_size=True)
                try:
                    info.index_signature = compute_project_index_signature(project_path)
                    info.index_cached_at = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    info.index_last_verified = info.index_cached_at
                    info.index_hit = False
                except Exception:
                    pass

                resolved = str(project_path.resolve())
                self.projects = [p for p in self.projects if str(p.path.resolve()) != resolved]
                self.projects.append(info)

                if add_to_locations:
                    parent = str(project_path.parent)
                    if parent not in self.scan_locations:
                        self.scan_locations.insert(0, parent)
                    self.root_path_var.set(parent)
                    self._refresh_location_combo()

                self._detect_duplicates()
                self._apply_filter(update_status=False)
                self._update_dashboard()
                self._update_status_totals()
                self._refresh_filter_values()
                self._save_project_index()

                iid = str(id(info))
                if iid in self.project_tree.get_children():
                    self.project_tree.selection_set(iid)
                    self.project_tree.see(iid)
                    self.selected_project = info
                    self._show_project_details(info)

                self.status_var.set(_tr('ui.source.nieuw.project.aangemaakt.p0.168c9112',p0=project_path))
                write_action_log("new_project", [
                    f"Project: {info.name}",
                    f"Type: {info.project_type}",
                    f"Pad: {project_path}",
                    "Bestanden:",
                    *[f"- {item}" for item in dialog.result.get("created_files", [])],
                ])

                if open_after:
                    windows_open_path(project_path)

            except Exception as exc:
                messagebox.showerror(_tr('ui.source.project.aangemaakt.analyse.mislukt.1fda63d9'), _tr('ui.source.projectmap.is.aangemaakt.p0.analyse.mislukt.p1.f37f88c2',p0=project_path,p1=exc))
        def _copy_output_to_clipboard(self) -> None:
            if not hasattr(self, "output_text"):
                return
            txt = self.output_text.get("1.0", END).strip()
            self.root.clipboard_clear()
            self.root.clipboard_append(txt)
            self.status_var.set(_tr('ui.source.uitvoer.gekopieerd.naar.klembord.44ba787c'))
        def _show_cleanup_items(self, p: ProjectInfo) -> None:
            for item in self.cleanup_tree.get_children():
                self.cleanup_tree.delete(item)

            for idx, item in enumerate(p.cleanup_items):
                self.cleanup_tree.insert(
                    "",
                    END,
                    iid=str(idx),
                    values=(
                        item.name,
                        item.risk,
                        item.category,
                        format_bytes(item.size_bytes),
                        item.file_count,
                        item.advice,
                        str(item.path),
                    ),
                )

            if p.cleanup_items:
                profile = self.cleanup_profile_var.get() if hasattr(self, "cleanup_profile_var") else "Veilig"
                profile_items = [item for item in p.cleanup_items if item.name in CLEANUP_PROFILES.get(profile, set())]
                profile_bytes = sum(item.size_bytes for item in profile_items)
                self.cleanup_summary_var.set(
                    _tr('ui.source.p0.opruimbare.map.pen.p1.totaal.profiel.p2.p3..0444fe88',p0=len(p.cleanup_items),p1=format_bytes(p.cleanup_bytes),p2=profile,p3=len(profile_items),p4=format_bytes(profile_bytes),p5=self._cleanup_risk_summary(profile_items or p.cleanup_items))
                )
                if hasattr(self, "cleanup_profile_summary_var"):
                    self.cleanup_profile_summary_var.set(cleanup_profile_summary(profile))
            else:
                self.cleanup_summary_var.set(_tr('ui.source.geen.bekende.build.cachemappen.gevonden.3ffefb77'))
                if hasattr(self, "cleanup_profile_summary_var"):
                    self.cleanup_profile_summary_var.set(cleanup_profile_summary(self.cleanup_profile_var.get()))
        def _find_project_backups(self, p: ProjectInfo) -> list[Path]:
            backup_dir = get_backup_dir()
            patterns = [
                f"{p.name}_*.zip",
                f"{p.name}_*.7z",
            ]
            found: list[Path] = []
            for pattern in patterns:
                found.extend(backup_dir.glob(pattern))
            return sorted(found, key=lambda x: x.stat().st_mtime if x.exists() else 0, reverse=True)
        def _show_backups_info(self, p: ProjectInfo) -> None:
            if not hasattr(self, "backups_text_var"):
                return
            backups = self._find_project_backups(p)
            lines = [f"Backupmap: {get_backup_dir()}", ""]
            if not backups:
                lines.append(_tr('ui.source.geen.backups.gevonden.voor.dit.project.085760a0'))
            else:
                for b in backups[:20]:
                    try:
                        size = format_bytes(b.stat().st_size)
                        modified = _dt.datetime.fromtimestamp(b.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
                    except Exception:
                        size = "?"
                        modified = "?"
                    lines.append(f"- {b.name} | {size} | {modified}")
                if len(backups) > 20:
                    lines.append(f"... plus {len(backups) - 20} extra backups.")
            self.backups_text_var.set("\n".join(lines))
        def _show_snapshot_info(self, p: ProjectInfo) -> None:
            if not hasattr(self, "snapshot_text_var"):
                return
            lines = [
                f"Snapshot aanwezig: {'Ja' if p.snapshot_present else 'Nee'}",
                f"Samenvatting: {p.snapshot_diff_summary or '-'}",
                "",
            ]
            lines.extend(p.snapshot_diff_details or [_tr('ui.source.maak.eerst.een.snapshot.d3f94546')])
            self.snapshot_text_var.set("\n".join(lines))
        def _export_requirements(self) -> None:
            p = self._require_selection()
            if not p:
                return
            self._open_terminal_at(p.path, "python -m pip freeze > requirements_export.txt")
            self.status_var.set(_tr('ui.source.requirements.export.txt.wordt.via.terminal.aan.6c92520c'))
        def _copy_project(self) -> None:
            p = self._require_selection()
            if not p:
                return

            dest_root = filedialog.askdirectory(title=_tr('ui.source.kies.doelmap.voor.kopie.46bfc171'))
            if not dest_root:
                return

            dest = Path(dest_root) / p.path.name
            if dest.exists():
                if not messagebox.askyesno(
                    _tr('ui.source.doel.bestaat.al.b2733f87'),
                    _tr('ui.source.deze.map.bestaat.al.p0.kopie.maken.met.suffix.4a2e0be5',p0=dest),
                ):
                    return
                counter = 2
                while True:
                    candidate = Path(dest_root) / f"{p.path.name}_copy{counter}"
                    if not candidate.exists():
                        dest = candidate
                        break
                    counter += 1

            try:
                self.status_var.set(_tr('ui.source.kopi.ren.naar.p0.27d9efe4',p0=dest))
                self.root.update_idletasks()
                shutil.copytree(p.path, dest, ignore=shutil.ignore_patterns(".git"))
                log = write_action_log("copy", [f"Project: {p.name}", f"Van: {p.path}", f"Naar: {dest}", "Git-map is niet meegekopieerd."])
                self.status_var.set(_tr('ui.source.kopie.gereed.p0.65494338',p0=dest))
                messagebox.showinfo(_tr('ui.source.kopie.gereed.c69ff684'), _tr('ui.source.project.gekopieerd.naar.p0.log.p1.01c5f725',p0=dest,p1=log))
            except Exception as exc:
                messagebox.showerror(_tr('ui.source.kopi.ren.mislukt.0b1c52dd'), str(exc))
                self.status_var.set(_tr('ui.source.kopi.ren.mislukt.7f4b7526'))
        def _move_project(self) -> None:
            """v6.0: veilig verplaatsen met ID-behoud, validatie en terugrolpoging."""
            p = self._require_selection()
            if not p:
                return
            self._ensure_project_id(p)
            dest_root = filedialog.askdirectory(title=_tr('ui.source.kies.doelmap.voor.veilig.verplaatsen.ef46fb49'))
            if not dest_root:
                return
            source = p.path.resolve()
            destination_root = Path(dest_root).resolve()
            dest = destination_root / source.name
            try:
                if destination_root == source or destination_root in source.parents:
                    messagebox.showerror(_tr('ui.source.ongeldig.doel.37678060'), _tr('ui.source.de.projectmap.kan.niet.naar.zichzelf.of.een.bo.4859d780'))
                    return
                if source in destination_root.parents:
                    messagebox.showerror(_tr('ui.source.ongeldig.doel.37678060'), _tr('ui.source.de.doelmap.mag.niet.binnen.het.project.liggen.1bd56407'))
                    return
            except Exception:
                pass
            if dest.exists():
                messagebox.showerror(_tr('ui.source.doel.bestaat.al.b2733f87'), _tr('ui.source.doelmap.bestaat.al.p0.cf29642c',p0=dest))
                return
            if not messagebox.askyesno(_tr('ui.source.safe.move.center.127e02d8'), _tr('ui.source.project.veilig.verplaatsen.project.id.p0.van.p.8581a661',p0=p.project_id,p1=source,p2=dest)):
                return
            old_path = source
            moved = False
            try:
                self.status_var.set(_tr('ui.source.veilig.verplaatsen.naar.p0.d20cb192',p0=dest))
                self.root.update_idletasks()
                destination_root.mkdir(parents=True, exist_ok=True)
                shutil.move(str(source), str(dest))
                moved = True
                if not dest.exists() or source.exists():
                    raise RuntimeError("Controle na verplaatsen is mislukt.")
                data = load_project_meta(dest)
                data["project_id"] = p.project_id
                data["last_moved"] = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                data["previous_path"] = str(old_path)
                save_project_meta(dest, data)
                check = load_project_meta(dest)
                if str(check.get("project_id", "")) != p.project_id:
                    raise RuntimeError("Project-ID kon na verplaatsen niet worden bevestigd.")
                p.path = dest
                p.name = dest.name
                p.last_moved = data["last_moved"]
                self._save_project_index()
                log = write_action_log("safe_move", [f"Project: {p.name}", f"Project-ID: {p.project_id}", f"Van: {old_path}", f"Naar: {dest}", "Validatie: geslaagd"])
                self._apply_filter(update_status=False)
                self._show_project_details(p)
                self.status_var.set(_tr('ui.source.veilig.verplaatst.p0.0184745d',p0=dest))
                messagebox.showinfo(_tr('ui.source.safe.move.gereed.da1ce492'), _tr('ui.source.project.veilig.verplaatst.naar.p0.project.id.b.22281044',p0=dest,p1=p.project_id,p2=log))
            except Exception as exc:
                rollback = "Niet nodig."
                if moved and dest.exists() and not old_path.exists():
                    try:
                        shutil.move(str(dest), str(old_path))
                        rollback = "Terugrol geslaagd."
                    except Exception as rb_exc:
                        rollback = f"Terugrol mislukt: {rb_exc}"
                write_action_log("safe_move_error", [f"Project: {p.name}", f"Van: {old_path}", f"Naar: {dest}", f"Fout: {exc}", rollback])
                self.status_var.set(_tr('ui.source.veilig.verplaatsen.mislukt.7c2fc391'))
                messagebox.showerror(_tr('ui.source.safe.move.mislukt.f647d2e4'), _tr('ui.source.p0.p1.11dc5127',p0=exc,p1=rollback))
        def _export_csv(self) -> None:
            if not self.projects:
                messagebox.showinfo(_tr('ui.source.geen.data.ec45aa59'), _tr('ui.source.er.zijn.nog.geen.projecten.gevonden.e1d933e5'))
                return

            default_name = f"projecten_{_dt.datetime.now().strftime('%Y-%m-%d')}.csv"
            path = filedialog.asksaveasfilename(
                title=_tr('ui.source.csv.exporteren.91893693'),
                defaultextension=".csv",
                initialfile=default_name,
                filetypes=[(_tr('ui.source.csv.32811883'), "*.csv"), (_tr('ui.source.alle.bestanden.3b98611e'), "*.*")],
            )

            if not path:
                return

            try:
                with open(path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f, delimiter=";")
                    writer.writerow(
                        [
                            "Naam",
                            "Type",
                            "Favoriet",
                            "Projectversie",
                            "Versierol",
                            "Groep",
                            "Startcommando",
                            "Build-readiness",
                            "Build-readiness score",
                            "Build-readiness redenen",
                            "Release-status",
                            "minSdk",
                            "targetSdk",
                            "compileSdk",
                            "Keystore genoemd",
                            "Embedded platform",
                            "Embedded board",
                            "Embedded framework",
                            "PlatformIO",
                            "PlatformIO envs",
                            "INO-bestand",
                            "GPIO",
                            "Camera",
                            "I2C/SPI",
                            "Systemd service",
                            "Shell scripts",
                            "Dependencies",
                            "Pad",
                            "Grootte bytes",
                            "Grootte",
                            "Laatst gewijzigd",
                            "Package Name",
                            "VersionCode",
                            "VersionName",
                            "React Native",
                            "Expo SDK",
                            "Python versie",
                            "Rust versie",
                            "Git",
                            "Git branch",
                            "Git remote",
                            "Gezondheid",
                            "Gezondheid score",
                            "Gebruiker status",
                            "Gebruiker tags",
                            "Notitie",
                            "APK",
                            "AAB",
                            "README",
                            "LICENSE",
                            "Tags",
                            "Opruimbaar bytes",
                            "Opruimbaar",
                            "Opruim mappen",
                            "Opruim risico",
                            "Duplicaat",
                        ]
                    )

                    for p in self.filtered_projects:
                        writer.writerow(
                            [
                                p.name,
                                p.project_type,
                                "Ja" if p.favorite else "Nee",
                                p.project_version_label,
                                p.version_role,
                                p.project_group,
                                p.start_command,
                                p.build_readiness_status,
                                p.build_readiness_score,
                                " | ".join(p.build_readiness_reasons),
                                p.android_release_status,
                                p.min_sdk,
                                p.target_sdk,
                                p.compile_sdk,
                                "Ja" if p.keystore_mentioned else "Nee",
                                p.embedded_platform,
                                p.embedded_board,
                                p.embedded_framework,
                                "Ja" if p.platformio_present else "Nee",
                                ", ".join(p.platformio_envs),
                                p.ino_file,
                                "Ja" if p.gpio_used else "Nee",
                                "Ja" if p.camera_used else "Nee",
                                "Ja" if p.i2c_spi_used else "Nee",
                                "Ja" if p.systemd_service_present else "Nee",
                                "Ja" if p.shell_scripts_present else "Nee",
                                len(p.dependencies),
                                str(p.path),
                                p.size_bytes,
                                format_bytes(p.size_bytes),
                                p.modified_text,
                                p.package_name,
                                p.version_code,
                                p.version_name,
                                p.react_native_version,
                                p.expo_sdk,
                                p.python_version,
                                p.rust_version,
                                "Ja" if p.git_present else "Nee",
                                p.git_branch,
                                p.git_remote,
                                p.health_status,
                                p.health_score,
                                p.user_status,
                                ", ".join(p.user_tags),
                                p.user_note.replace("\n", " "),
                                "Ja" if p.apk_present else "Nee",
                                "Ja" if p.aab_present else "Nee",
                                "Ja" if p.readme_present else "Nee",
                                "Ja" if p.license_present else "Nee",
                                ", ".join(p.tags),
                                p.cleanup_bytes,
                                format_bytes(p.cleanup_bytes),
                                len(p.cleanup_items),
                                self._cleanup_risk_summary(p.cleanup_items) if p.cleanup_items else "-",
                                p.duplicate_hint,
                            ]
                        )

                self.status_var.set(_tr('ui.source.csv.ge.xporteerd.p0.914e1d3d',p0=path))
                messagebox.showinfo(_tr('ui.source.csv.gereed.86d3e19c'), _tr('ui.source.csv.opgeslagen.p0.93a953eb',p0=path))
            except Exception as exc:
                messagebox.showerror(_tr('ui.source.csv.export.mislukt.0ad85ee9'), str(exc))
        def _update_cleanup_profile_summary(self) -> None:
            """Werk de uitleg bij het geselecteerde cleanup-profiel bij."""
            if hasattr(self, "cleanup_profile_summary_var"):
                self.cleanup_profile_summary_var.set(cleanup_profile_summary(self.cleanup_profile_var.get()))
            if self.selected_project:
                self._show_cleanup_items(self.selected_project)
        def _make_backup_for_cleanup(self, p: ProjectInfo) -> Path | None:
            """Maak een ZIP-backup voorafgaand aan opschonen."""
            if not bool(self.backup_before_cleanup_var.get()):
                return None

            backup_path = self.archiver.archive_project(
                project=p,
                destination_dir=get_backup_dir(),
                use_7z=False,
                excludes={"node_modules": False, "build": False, "dist": False, ".gradle": False, "target": False, ".pio": False},
            )
            stamp = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            try:
                update_project_meta(p.path, {"last_backup": stamp})
                p.last_backup = stamp
            except Exception:
                pass
            write_action_log("cleanup_backup", [f"Project: {p.name}", f"Backup: {backup_path}"])
            return backup_path
        def _cleanup_risk_summary(self, items: list[CleanupItem]) -> str:
            """Maak korte risicosamenvatting."""
            counts = {"Laag": 0, "Middel": 0, "Hoog": 0}
            sizes = {"Laag": 0, "Middel": 0, "Hoog": 0}
            for item in items:
                counts[item.risk] = counts.get(item.risk, 0) + 1
                sizes[item.risk] = sizes.get(item.risk, 0) + item.size_bytes
            parts = []
            for risk in ["Hoog", "Middel", "Laag"]:
                if counts.get(risk):
                    parts.append(f"{risk}: {counts[risk]} map(pen), {format_bytes(sizes[risk])}")
            return "; ".join(parts) if parts else "Geen opruimitems."
        def _show_cleanup_wizard(self) -> None:
            """Toon een veilige opruimwizard met profieladvies."""
            p = self._require_selection()
            if not p:
                return
            if not p.cleanup_items:
                messagebox.showinfo(_tr('ui.source.opruimwizard.40b0a095'), _tr('ui.source.geen.bekende.build.cachemappen.gevonden.3ffefb77'))
                return

            lines = [
                "OPRUIMWIZARD",
                "",
                _tr('ui.source.project.p0.e06d67de',p0=p.name),
                f"Pad: {p.path}",
                "",
                _tr('ui.source.doel.f8ba2dc9'),
                _tr('ui.source.ruimte.vrijmaken.zonder.broncode.te.verwijdere.587bee99'),
                "",
                "Profielen:",
            ]

            for profile, names in CLEANUP_PROFILES.items():
                items = [item for item in p.cleanup_items if item.name in names]
                total = sum(item.size_bytes for item in items)
                high = sum(1 for item in items if item.risk == "Hoog")
                lines.append(f"- {profile}: {len(items)} map(pen), {format_bytes(total)}")
                lines.append(f"  {cleanup_profile_summary(profile)}")
                if high:
                    lines.append(_tr('ui.source.let.op.bevat.p0.hoog.risico.dependency.omgevin.ecd2a30f',p0=high))
                lines.append("")

            lines.append(_tr('ui.source.gevonden.opruimitems.b0043b11'))
            for item in p.cleanup_items:
                lines.append(
                    _tr('ui.source.p0.p1.p2.p3.p4.bestanden.cff5de08',p0=item.risk,p1=item.name,p2=item.category,p3=format_bytes(item.size_bytes),p4=item.file_count)
                )
                lines.append(f"  {item.advice}")
                lines.append(f"  {item.path}")

            lines.extend([
                "",
                "Advies:",
                "- Begin met profiel 'Veilig'.",
                "- Gebruik 'Build-output' als buildmappen veel ruimte innemen.",
                _tr('ui.source.gebruik.dependencies.of.grondig.alleen.als.opn.1a2b6d09'),
                _tr('ui.source.laat.backup.v.r.opschonen.aan.staan.597dc2eb'),
            ])

            self._show_text_window(_tr('ui.source.opruimwizard.40b0a095'), "\n".join(lines))
        def _cleanup_items_for_current_profile(self, p: ProjectInfo) -> list[CleanupItem]:
            """Geef opruimitems op basis van geselecteerde cleanup-profiel."""
            selected_ids = self.cleanup_tree.selection() if hasattr(self, "cleanup_tree") else []
            if selected_ids:
                return [p.cleanup_items[int(iid)] for iid in selected_ids]
            profile = self.cleanup_profile_var.get() if hasattr(self, "cleanup_profile_var") else "Veilig"
            allowed = CLEANUP_PROFILES.get(profile, CLEANUP_PROFILES["Veilig"])
            return [item for item in p.cleanup_items if item.name in allowed]
        def _cleanup_selected_project(self, dry_run: bool = False) -> None:
            p = self._require_selection()
            if not p:
                return

            if not p.cleanup_items:
                messagebox.showinfo(_tr('ui.source.geen.opruimactie.37c80378'), _tr('ui.source.geen.bekende.build.cachemappen.gevonden.3ffefb77'))
                return

            items = self._cleanup_items_for_current_profile(p)
            if not items:
                messagebox.showinfo(_tr('ui.source.geen.opruimactie.37c80378'), _tr('ui.source.geen.opruimitems.binnen.profiel.p0.83d2bc13',p0=self.cleanup_profile_var.get()))
                return

            total_bytes = sum(i.size_bytes for i in items)
            high_risk = [i for i in items if i.risk == "Hoog"]
            backup_path: Path | None = None

            lines = [
                "DRY-RUN: de volgende mappen zouden worden verwijderd:" if dry_run else "De volgende mappen worden verwijderd:",
                "",
                _tr('ui.source.project.p0.e06d67de',p0=p.name),
                f"Profiel: {self.cleanup_profile_var.get()}",
                _tr('ui.source.risico.p0.28ace877',p0=self._cleanup_risk_summary(items)),
                "",
            ]
            for item in items:
                lines.append(f"- [{item.risk}] {item.path} ({format_bytes(item.size_bytes)})")
                lines.append(f"  {item.advice}")

            lines.extend(
                [
                    "",
                    f"Geschatte ruimtebesparing: {format_bytes(total_bytes)}",
                    "",
                    _tr('ui.source.broncode.wordt.niet.bewust.verwijderd.maar.con.d5fddc71'),
                ]
            )

            if high_risk:
                lines.extend([
                    "",
                    _tr('ui.source.let.op.dit.profiel.bevat.hoog.risico.dependenc.28237062'),
                    "Voorbeelden: node_modules, .venv, venv of env.",
                    _tr('ui.source.na.verwijderen.kan.opnieuw.installeren.nodig.z.e02c0d0b'),
                ])

            log_lines = [
                f"Project: {p.name}",
                f"Pad: {p.path}",
                f"Profiel: {self.cleanup_profile_var.get()}",
                f"Dry-run: {'ja' if dry_run else 'nee'}",
                f"Backup vóór opschonen: {'ja' if self.backup_before_cleanup_var.get() else 'nee'}",
                f"Risico: {self._cleanup_risk_summary(items)}",
                f"Geschatte ruimtebesparing: {format_bytes(total_bytes)}",
                "",
                "Mappen:",
            ] + [
                f"- [{item.risk}] {item.path} ({format_bytes(item.size_bytes)}, {item.file_count} bestanden) | {item.advice}"
                for item in items
            ]

            log_path = write_action_log("cleanup_dry_run" if dry_run else "cleanup_plan", log_lines)

            if dry_run:
                messagebox.showinfo(_tr('ui.source.dry.run.opschonen.0d2b895f'), "\n".join(lines) + f"\n\nLog:\n{log_path}")
                self.status_var.set(_tr('ui.source.dry.run.gereed.log.p0.b9109ab2',p0=log_path))
                return

            if high_risk and not self.backup_before_cleanup_var.get():
                if not messagebox.askyesno(
                    _tr('ui.source.hoog.risico.zonder.backup.a3cafe0c'),
                    _tr('ui.source.er.staan.hoog.risico.dependency.omgevingmappen.26d59fa9'),
                ):
                    return

            if not messagebox.askyesno(_tr('ui.source.opschonen.bevestigen.fc167c37'), "\n".join(lines) + _tr('ui.source.log.wordt.bijgewerkt.p0.0ec297fd',p0=log_path)):
                return

            if self.backup_before_cleanup_var.get():
                try:
                    self.status_var.set(_tr('ui.source.backup.maken.v.r.opschonen.1580ac30'))
                    self.root.update_idletasks()
                    backup_path = self._make_backup_for_cleanup(p)
                except Exception as exc:
                    messagebox.showerror(_tr('ui.source.backup.v.r.opschonen.mislukt.31a9af9d'), _tr('ui.source.opschonen.is.gestopt.omdat.de.backup.niet.lukt.a11cbf91',p0=exc))
                    self.status_var.set(_tr('ui.source.opschonen.gestopt.backup.mislukt.82845a90'))
                    return

            # Tweede bevestiging bewust kort en hard.
            if not messagebox.askyesno(
                _tr('ui.source.laatste.bevestiging.00d224af'),
                _tr('ui.source.zeker.weten.dat.deze.mappen.verwijderd.mogen.w.ffeafb08'),
            ):
                return

            removed = 0
            failed = []

            for item in items:
                try:
                    if item.path.exists() and item.path.is_dir():
                        shutil.rmtree(item.path)
                        removed += 1
                except Exception as exc:
                    failed.append(f"{item.path}: {exc}")

            stamp_cleanup = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            try:
                update_project_meta(p.path, {"last_cleanup": stamp_cleanup})
                p.last_cleanup = stamp_cleanup
            except Exception:
                pass

            result_lines = log_lines + ["", f"Backup: {backup_path or '-'}", f"Verwijderd: {removed}"]
            if failed:
                result_lines.append("Niet verwijderd:")
                result_lines.extend(failed)
            result_log = write_action_log("cleanup_result", result_lines)

            if failed:
                messagebox.showwarning(
                    _tr('ui.source.opschonen.deels.mislukt.1e820301'),
                    f"{removed} map(pen) verwijderd.\n\nNiet verwijderd:\n" + "\n".join(failed[:10]) + f"\n\nLog:\n{result_log}",
                )
            else:
                messagebox.showinfo(
                    _tr('ui.source.opschonen.gereed.a6e70e64'),
                    _tr('ui.source.p0.map.pen.verwijderd.backup.p1.log.p2.e955fc22',p0=removed,p1=backup_path or '-',p2=result_log),
                )

            self._refresh_cleanup_analysis()
        def _backup_project_for_delete(self, project: ProjectInfo) -> Path | None:
            """Maak een volledige ZIP-backup voordat een project verwijderd wordt."""
            backup_path = self.archiver.archive_project(
                project=project,
                destination_dir=get_backup_dir(),
                use_7z=False,
                excludes={
                    "node_modules": False,
                    "build": False,
                    "dist": False,
                    ".gradle": False,
                    "target": False,
                    ".pio": False,
                    ".venv": False,
                    "venv": False,
                    "env": False,
                },
            )
            return backup_path
        def _delete_selected_project(self) -> None:
            """Veilige verwijderwizard voor één project."""
            p = self._require_selection()
            if not p:
                return

            if not p.path.exists():
                if messagebox.askyesno(
                    _tr('ui.source.projectmap.bestaat.niet.7927b200'),
                    _tr('ui.source.de.projectmap.bestaat.niet.meer.p0.alleen.uit..34eb9cb8',p0=p.path),
                ):
                    self._remove_project_from_memory_and_index(p)
                    self.status_var.set(_tr('ui.source.project.uit.index.verwijderd.11e957c1'))
                return

            if normalize_workflow_status(p.user_status) != "Weggooien?":
                if messagebox.askyesno(
                    _tr('ui.source.status.vereist.02e0b398'),
                    _tr('ui.source.voor.veilig.verwijderen.moet.de.workflowstatus.bd388c7a'),
                ):
                    self._mark_project_as_disposable_and_stop(p)
                return

            dialog = DeleteProjectDialog(self.root, p)
            self.root.wait_window(dialog.window)

            if not dialog.result:
                return

            make_backup = bool(dialog.result.get("backup", True))
            use_recycle = bool(dialog.result.get("recycle", True))

            if not messagebox.askyesno(
                _tr('ui.source.laatste.bevestiging.00d224af'),
                _tr('ui.source.laatste.controle.project.p0.pad.p1.backup.make.1fdebef0',p0=p.name,p1=p.path,p2='ja' if make_backup else 'nee',p3='ja' if use_recycle else 'nee'),
            ):
                return

            backup_path = None
            log_lines = [
                _tr('ui.source.project.p0.e06d67de',p0=p.name),
                f"Pad: {p.path}",
                f"Workflowstatus: {p.user_status}",
                f"Backup gevraagd: {'ja' if make_backup else 'nee'}",
                _tr('ui.source.prullenbak.gevraagd.p0.961b45fc',p0='ja' if use_recycle else 'nee'),
                f"Grootte: {format_bytes(p.size_bytes)}",
                f"Laatst gewijzigd: {p.modified_text or '-'}",
            ]

            try:
                if make_backup:
                    self.status_var.set(_tr('ui.source.backup.maken.v.r.verwijderen.81a1323a'))
                    self.root.update_idletasks()
                    backup_path = self._backup_project_for_delete(p)
                    log_lines.append(f"Backup: {backup_path}")

                self.status_var.set(_tr('ui.source.project.verwijderen.ab9226ed'))
                self.root.update_idletasks()

                if use_recycle:
                    success, message = move_path_to_recycle_bin(p.path)
                    method = "Prullenbak"
                else:
                    success, message = permanently_delete_path(p.path)
                    method = "Permanent"

                if not success and use_recycle:
                    log_lines.append(_tr('ui.source.prullenbak.mislukt.p0.4c9740b5',p0=message))
                    if messagebox.askyesno(
                        _tr('ui.source.prullenbak.mislukt.efde3b7a'),
                        _tr('ui.source.verplaatsen.naar.de.prullenbak.is.mislukt.p0.p.a75313a3',p0=message),
                    ):
                        success, message = permanently_delete_path(p.path)
                        method = "Permanent na prullenbak-fout"

                log_lines.extend([
                    f"Methode: {method}",
                    _tr('ui.source.resultaat.p0.be4c0fd4',p0='gelukt' if success else 'mislukt'),
                    f"Melding: {message}",
                ])
                log = write_action_log("project_delete", log_lines)

                if success:
                    self._remove_project_from_memory_and_index(p)
                    self.status_var.set(_tr('ui.source.project.verwijderd.p0.33f33bd4',p0=p.name))
                    messagebox.showinfo(
                        _tr('ui.source.project.verwijderd.4587680d'),
                        _tr('ui.source.project.verwijderd.p0.methode.p1.backup.p2.log.acba61cb',p0=p.name,p1=method,p2=backup_path or '-',p3=log),
                    )
                else:
                    self.status_var.set(_tr('ui.source.project.verwijderen.mislukt.f7dbcac8'))
                    messagebox.showerror(
                        _tr('ui.source.verwijderen.mislukt.3533d122'),
                        _tr('ui.source.project.is.niet.verwijderd.p0.log.p1.2ccbfc23',p0=message,p1=log),
                    )

            except Exception as exc:
                log_lines.extend(["Exception:", str(exc), traceback.format_exc()])
                log = write_action_log("project_delete_error", log_lines)
                messagebox.showerror(_tr('ui.source.verwijderen.mislukt.3533d122'), _tr('ui.source.p0.log.p1.791d4453',p0=exc,p1=log))
                self.status_var.set(_tr('ui.source.project.verwijderen.mislukt.f7dbcac8'))
        def _bulk_delete_selected_disposable(self) -> None:
            """Bulkverwijdering voor geselecteerde projecten met status Weggooien?."""
            projects = self._get_selected_projects()
            if not projects:
                messagebox.showinfo(_tr('ui.source.geen.selectie.60b466be'), _tr('ui.source.selecteer.eerst.n.of.meer.projecten.9b241af5'))
                return

            eligible = [p for p in projects if normalize_workflow_status(p.user_status) == "Weggooien?"]
            skipped = [p for p in projects if normalize_workflow_status(p.user_status) != "Weggooien?"]

            if not eligible:
                messagebox.showinfo(
                    _tr('ui.source.geen.verwijderbare.selectie.3f05a90e'),
                    _tr('ui.source.geen.geselecteerd.project.heeft.workflowstatus.8f6dd3d4'),
                )
                return

            lines = [
                _tr('ui.source.te.verwijderen.via.prullenbak.p0.project.en.99f865d5',p0=len(eligible)),
                f"Overgeslagen zonder status Weggooien?: {len(skipped)}",
                "",
                "Projecten:",
            ]
            lines.extend(f"- {p.name}\n  {p.path}" for p in eligible)
            lines.extend([
                "",
                _tr('ui.source.voor.elk.project.wordt.eerst.een.zip.backup.ge.b73a73f0'),
                _tr('ui.source.projecten.worden.na.succesvolle.verwijdering.u.c01b4071'),
            ])

            if not messagebox.askyesno(_tr('ui.source.bulk.verwijderen.bevestigen.3b445edd'), "\n".join(lines)):
                return

            removed: list[ProjectInfo] = []
            failed: list[str] = []
            logs: list[str] = []

            for p in eligible:
                if not p.path.exists():
                    removed.append(p)
                    logs.append(f"{p.name}: map bestond niet meer, alleen uit index verwijderd.")
                    continue

                try:
                    backup_path = self._backup_project_for_delete(p)
                    success, message = move_path_to_recycle_bin(p.path)
                    if success:
                        removed.append(p)
                        logs.append(f"{p.name}: verwijderd via prullenbak | backup: {backup_path}")
                    else:
                        failed.append(f"{p.name}: {message}")
                except Exception as exc:
                    failed.append(f"{p.name}: {exc}")

            for p in removed:
                try:
                    self._remove_project_from_memory_and_index(p)
                except Exception:
                    pass

            log = write_action_log(
                "bulk_project_delete",
                [
                    f"Gevraagd: {len(projects)}",
                    f"Eligible Weggooien?: {len(eligible)}",
                    f"Verwijderd: {len(removed)}",
                    f"Mislukt: {len(failed)}",
                    "",
                    "Resultaten:",
                    *logs,
                    "",
                    "Fouten:",
                    *failed,
                ],
            )

            self._detect_duplicates()
            self._apply_filter(update_status=False)
            self._refresh_filter_values()
            self._update_dashboard()
            self._update_status_totals()
            self._save_project_index()

            if failed:
                messagebox.showwarning(
                    _tr('ui.source.bulk.verwijderen.deels.mislukt.9d4d3b0b'),
                    _tr('ui.source.verwijderd.p0.mislukt.p1.log.p2.8cb4be93',p0=len(removed),p1=len(failed),p2=log),
                )
            else:
                messagebox.showinfo(
                    _tr('ui.source.bulk.verwijderen.gereed.8ab45821'),
                    _tr('ui.source.verwijderd.p0.project.en.log.p1.8eef006e',p0=len(removed),p1=log),
                )
        def _archive_selected_project(self) -> None:
            p = self._require_selection()
            if not p:
                return

            destination = filedialog.askdirectory(title=_tr('ui.source.kies.map.voor.archiefbestand.f630aeec'))
            if not destination:
                return

            dialog = ArchiveOptionsDialog(self.root, self.archiver.seven_zip)
            self.root.wait_window(dialog.window)

            if not dialog.result:
                return

            use_7z = dialog.result["use_7z"]
            excludes = dialog.result["excludes"]

            try:
                self.status_var.set(_tr('ui.source.archiveren.gestart.9aae99a2'))
                self.root.update_idletasks()

                def progress(done, total):
                    self.status_var.set(_tr('ui.source.archiveren.p0.p1.ea76b383',p0=done,p1=total))
                    self.root.update_idletasks()

                archive_path = self.archiver.archive_project(
                    project=p,
                    destination_dir=Path(destination),
                    use_7z=use_7z,
                    excludes=excludes,
                    progress_callback=progress,
                )
                stamp_archive = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                try:
                    update_project_meta(p.path, {"last_archive": stamp_archive})
                    p.last_archive = stamp_archive
                except Exception:
                    pass
                log = write_action_log(
                    "archive",
                    [
                        f"Project: {p.name}",
                        f"Bron: {p.path}",
                        f"Archief: {archive_path}",
                        f"Formaat: {'7Z' if use_7z else 'ZIP'}",
                        "Excludes:",
                        *[f"- {name}: {'ja' if active else 'nee'}" for name, active in excludes.items()],
                    ],
                )

                self.status_var.set(_tr('ui.source.archief.gereed.p0.314978df',p0=archive_path))
                messagebox.showinfo(_tr('ui.source.archief.gereed.6565fea0'), _tr('ui.source.archief.gemaakt.p0.log.p1.9b681f44',p0=archive_path,p1=log))

            except subprocess.CalledProcessError as exc:
                messagebox.showerror(
                    _tr('ui.source.7z.archiveren.mislukt.2b9267e5'),
                    _tr('ui.source.7.zip.gaf.een.foutmelding.p0.18ed1ce3',p0=exc.stderr or exc),
                )
                self.status_var.set(_tr('ui.source.archiveren.mislukt.066529d3'))
            except Exception as exc:
                messagebox.showerror(_tr('ui.source.archiveren.mislukt.1a0115f2'), str(exc))
                self.status_var.set(_tr('ui.source.archiveren.mislukt.066529d3'))
        def _open_backup_folder(self) -> None:
            windows_open_path(get_backup_dir())
        def _backup_selected_project(self) -> None:
            p = self._require_selection()
            if not p:
                return
            try:
                self.status_var.set(_tr('ui.source.backup.maken.f3c45d87'))
                self.root.update_idletasks()
                backup_path = self.archiver.archive_project(
                    project=p,
                    destination_dir=get_backup_dir(),
                    use_7z=False,
                    excludes={"node_modules": False, "build": False, "dist": False, ".gradle": False, "target": False},
                )
                stamp = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                update_project_meta(p.path, {"last_backup": stamp})
                p.last_backup = stamp
                self._show_backups_info(p)
                log = write_action_log("backup", [f"Project: {p.name}", f"Backup: {backup_path}"])
                self.status_var.set(_tr('ui.source.backup.gereed.p0.f30c57d6',p0=backup_path))
                messagebox.showinfo(_tr('ui.source.backup.gereed.e770139c'), _tr('ui.source.backup.gemaakt.p0.log.p1.9dd44810',p0=backup_path,p1=log))
            except Exception as exc:
                messagebox.showerror(_tr('ui.source.backup.mislukt.2ca4a044'), str(exc))
                self.status_var.set(_tr('ui.source.backup.mislukt.8b4af763'))
        def _create_project_snapshot(self) -> None:
            p = self._require_selection()
            if not p:
                return
            target = p.path / SNAPSHOT_FILE
            snapshot = build_project_snapshot(p.path)
            target.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False), encoding="utf-8")
            log = write_action_log("snapshot", [f"Project: {p.name}", f"Snapshot: {target}", f"Bestanden: {snapshot.get('file_count')}"])
            self.status_var.set(_tr('ui.source.snapshot.gemaakt.log.p0.05333019',p0=log))
            self._refresh_selected_after_file_change()
        def _open_backup_dir(self) -> None:
            windows_open_path(get_backup_dir())
        def _bulk_export_csv_selected(self) -> None:
            projects = self._get_selected_projects()
            if not projects:
                messagebox.showinfo(_tr('ui.source.geen.selectie.60b466be'), _tr('ui.source.selecteer.eerst.n.of.meer.projecten.9b241af5'))
                return
            default_name = f"projecten_selectie_{_dt.datetime.now().strftime('%Y-%m-%d')}.csv"
            path = filedialog.asksaveasfilename(
                title=_tr('ui.source.csv.van.selectie.opslaan.77af1871'),
                defaultextension=".csv",
                initialfile=default_name,
                filetypes=[(_tr('ui.source.csv.32811883'), "*.csv"), (_tr('ui.source.alle.bestanden.3b98611e'), "*.*")],
            )
            if not path:
                return
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f, delimiter=";")
                writer.writerow(["Naam", "Workflow", "Type", "Groep", "Kwaliteit", "Build", "Grootte", "Opruimbaar", "Pad"])
                for p in projects:
                    writer.writerow([p.name, p.user_status, p.project_type, p.project_group, p.quality_status, p.build_readiness_status, format_bytes(p.size_bytes), format_bytes(p.cleanup_bytes), str(p.path)])
            self.status_var.set(_tr('ui.source.csv.selectie.opgeslagen.p0.4df186ec',p0=path))
        def _bulk_cleanup_dry_run(self) -> None:
            projects = self._get_selected_projects()
            if not projects:
                messagebox.showinfo(_tr('ui.source.geen.selectie.60b466be'), _tr('ui.source.selecteer.eerst.n.of.meer.projecten.9b241af5'))
                return
            lines = []
            total = 0
            for p in projects:
                items = self._cleanup_items_for_current_profile(p)
                subtotal = sum(i.size_bytes for i in items)
                total += subtotal
                lines.append(f"{p.name} - {format_bytes(subtotal)}")
                for item in items:
                    lines.append(f"  - {item.path} ({format_bytes(item.size_bytes)})")
                lines.append("")
            lines.insert(0, f"Bulk dry-run cleanup - {len(projects)} project(en)")
            lines.insert(1, f"Geschatte ruimtebesparing totaal: {format_bytes(total)}")
            lines.insert(2, "")
            log = write_action_log("bulk_cleanup_dry_run", lines)
            self._show_text_window(_tr('ui.source.bulk.dry.run.cleanup.17172c87'), "\n".join(lines) + f"\nLog: {log}")
        def _bulk_archive_selected(self) -> None:
            projects = self._get_selected_projects()
            if not projects:
                messagebox.showinfo(_tr('ui.source.geen.selectie.60b466be'), _tr('ui.source.selecteer.eerst.n.of.meer.projecten.9b241af5'))
                return
            destination = filedialog.askdirectory(title=_tr('ui.source.kies.map.voor.bulkarchieven.8ba5dd96'))
            if not destination:
                return
            dialog = ArchiveOptionsDialog(self.root, self.archiver.seven_zip)
            self.root.wait_window(dialog.window)
            if not dialog.result:
                return
            use_7z = dialog.result["use_7z"]
            excludes = dialog.result["excludes"]
            made = []
            for p in projects:
                try:
                    made.append(str(self.archiver.archive_project(p, Path(destination), use_7z, excludes)))
                except Exception as exc:
                    made.append(f"MISLUKT {p.name}: {exc}")
            log = write_action_log("bulk_archive", ["Archieven:", *made])
            self.status_var.set(_tr('ui.source.bulkarchief.gereed.log.p0.47a071d4',p0=log))
            self._show_text_window(_tr('ui.source.bulkarchief.03a49c86'), "\n".join(made) + f"\n\nLog: {log}")
        def _repair_backup_project(self, project: ProjectInfo) -> Path:
            """Maak vóór herstel een volledige ZIP-back-up, zonder zware cachemappen."""
            stamp = _dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
            safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", project.name).strip("_") or "project"
            path = get_backup_dir() / f"repair_{safe}_{stamp}.zip"
            excluded = {".git", "node_modules", "build", "dist", "target", ".gradle", ".pio", ".venv", "venv", "__pycache__"}
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as zf:
                for dirpath, dirnames, filenames in os.walk(project.path):
                    dirnames[:] = [d for d in dirnames if d not in excluded]
                    for filename in filenames:
                        src = Path(dirpath) / filename
                        try:
                            zf.write(src, src.relative_to(project.path))
                        except Exception:
                            pass
            return path
        def _export_project_library(self) -> None:
            if not self.projects:
                messagebox.showinfo(_tr('ui.source.geen.projecten.b183d12f'), _tr('ui.source.er.zijn.geen.projecten.om.te.exporteren.c4fad753'))
                return
            path = filedialog.asksaveasfilename(title=_tr('ui.source.portable.projectbibliotheek.exporteren.3a71ef48'), defaultextension=".json", initialfile=PROJECT_LIBRARY_FILE_NAME, filetypes=[(_tr('ui.source.json.031a4e76'), "*.json")])
            if not path:
                return
            try:
                data = self._portable_library_data()
                write_json_file(Path(path), data)
                log = write_action_log("library_export", [f"Bestand: {path}", f"Projecten: {len(self.projects)}"])
                self.status_var.set(_tr('ui.source.projectbibliotheek.ge.xporteerd.p0.f4c88106',p0=path))
                messagebox.showinfo(_tr('ui.source.export.gereed.08ecb02d'), _tr('ui.source.portable.bibliotheek.opgeslagen.p0.projecten.p.a390db27',p0=path,p1=len(self.projects),p2=log))
            except Exception as exc:
                messagebox.showerror(_tr('ui.source.export.mislukt.64f5bcf1'), str(exc))
        def _import_project_library(self) -> None:
            path = filedialog.askopenfilename(title=_tr('ui.source.portable.projectbibliotheek.importeren.6af7f160'), filetypes=[(_tr('ui.source.project.library.json.87a1bdc5'), "*.json"), (_tr('ui.source.alle.bestanden.3b98611e'), "*.*")])
            if not path:
                return
            data = load_json_file(Path(path), {})
            records = data.get("projects", []) if isinstance(data, dict) else []
            if data.get("schema") != "project-manager-portable-library" or not isinstance(records, list):
                messagebox.showerror(_tr('ui.source.ongeldige.bibliotheek.36996ec3'), _tr('ui.source.dit.bestand.is.geen.geldige.portable.projectbi.df6b3d36'))
                return
            by_id = {self._ensure_project_id(p): p for p in self.projects}
            by_path = {str(p.path).lower(): p for p in self.projects}
            updated = added = missing = 0
            for record in records:
                if not isinstance(record, dict):
                    continue
                rid = str(record.get("project_id", "")).strip()
                rpath = Path(str(record.get("path", "")))
                target = by_id.get(rid) or by_path.get(str(rpath).lower())
                if target:
                    target.project_id = rid or self._ensure_project_id(target)
                    target.user_status = normalize_workflow_status(record.get("workflow_status", target.user_status))
                    target.favorite = bool(record.get("favorite", target.favorite))
                    target.project_group = str(record.get("group", target.project_group))
                    target.user_note = str(record.get("note", target.user_note))
                    tags = record.get("user_tags", target.user_tags)
                    if isinstance(tags, list): target.user_tags = [str(x) for x in tags]
                    update_project_meta(target.path, {"project_id": target.project_id, "status": target.user_status, "favorite": target.favorite, "project_group": target.project_group, "notitie": target.user_note, "tags": target.user_tags})
                    updated += 1
                elif rpath.exists() and rpath.is_dir():
                    try:
                        project = self.analyzer.analyze(rpath, deep_size=False, fast_scan=True)
                        if rid:
                            project.project_id = rid
                            update_project_meta(rpath, {"project_id": rid})
                        self.projects.append(project)
                        by_id[project.project_id] = project
                        by_path[str(project.path).lower()] = project
                        added += 1
                    except Exception:
                        missing += 1
                else:
                    missing += 1
            self._analyze_all_project_intelligence()
            self._save_project_index()
            self._apply_filter(update_status=False)
            log = write_action_log("library_import", [f"Bestand: {path}", f"Bijgewerkt: {updated}", f"Toegevoegd: {added}", f"Niet gevonden: {missing}"])
            messagebox.showinfo(_tr('ui.source.import.gereed.d471afe1'), _tr('ui.source.bijgewerkt.p0.toegevoegd.p1.paden.niet.gevonde.ce434b77',p0=updated,p1=added,p2=missing,p3=log))
        def _backup_library_metadata(self) -> Path:
            stamp=_dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
            path=get_backup_dir()/f"project_library_metadata_{stamp}.json"
            write_json_file(path,self._portable_library_data())
            return path
        def _backup_selected_projects_v6(self) -> Path:
            projects=self._get_selected_projects() or ([self.selected_project] if self.selected_project else [])
            projects=[p for p in projects if p]
            if not projects: raise RuntimeError("Selecteer eerst één of meer projecten.")
            stamp=_dt.datetime.now().strftime("%Y-%m-%d_%H%M%S"); path=get_backup_dir()/f"project_selection_{stamp}.zip"
            with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED,allowZip64=True) as zf:
                zf.writestr(PROJECT_LIBRARY_FILE_NAME,json.dumps(self._portable_library_data(projects),ensure_ascii=False,indent=2))
                for p in projects:
                    for dirpath,dirnames,filenames in os.walk(p.path):
                        dirnames[:]=[d for d in dirnames if d not in {".git","node_modules","build","dist","target",".gradle",".pio",".venv","venv"}]
                        for filename in filenames:
                            src=Path(dirpath)/filename
                            try: zf.write(src,Path(p.project_id or p.name)/src.relative_to(p.path))
                            except Exception: pass
            return path
        def _show_backup_center(self) -> None:
            win=self._new_tool_window(); win.title(_tr('ui.source.back.upcentrum.v6.0.95c6f08b')); win.geometry("720x470")
            frame=ttk.Frame(win,padding=16); frame.pack(fill=BOTH,expand=True)
            ttk.Label(frame,text=_tr('ui.source.back.upcentrum.12e6edc8'),style="Title.TLabel").pack(anchor="w")
            ttk.Label(frame,text=_tr('ui.source.back.upmap.p0.4c51b12a',p0=get_backup_dir()),style="Muted.TLabel",wraplength=660).pack(anchor="w",pady=(0,14))
            result=StringVar(value=_tr('ui.source.kies.een.back.uptype.6072ac1c'))
            def run(kind):
                try:
                    if kind=="metadata": path=self._backup_library_metadata()
                    elif kind=="selection": path=self._backup_selected_projects_v6()
                    else:
                        stamp=_dt.datetime.now().strftime("%Y-%m-%d_%H%M%S"); path=get_backup_dir()/f"portfolio_full_{stamp}.zip"
                        with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED,allowZip64=True) as zf:
                            zf.writestr(PROJECT_LIBRARY_FILE_NAME,json.dumps(self._portable_library_data(),ensure_ascii=False,indent=2))
                            for p in self.projects:
                                for dirpath,dirnames,filenames in os.walk(p.path):
                                    dirnames[:]=[d for d in dirnames if d not in {".git","node_modules","build","dist","target",".gradle",".pio",".venv","venv"}]
                                    for filename in filenames:
                                        src=Path(dirpath)/filename
                                        try: zf.write(src,Path(self._ensure_project_id(p))/src.relative_to(p.path))
                                        except Exception: pass
                    result.set(f"Gereed: {path}"); write_action_log("backup_center",[f"Type: {kind}",f"Bestand: {path}"])
                except Exception as exc: messagebox.showerror(_tr('ui.source.back.up.mislukt.d7eea428'),str(exc))
            for text,kind,desc in [("Alleen bibliotheekmetadata","metadata","Snel; portable project_library.json met alle portfolio-informatie."),("Geselecteerde projecten","selection","ZIP met metadata en geselecteerde bronbestanden; zware build/dependencymappen uitgesloten."),("Volledig portfolio","full","ZIP met metadata en alle gevonden projecten; zware build/dependencymappen uitgesloten.")]:
                card=ttk.Frame(frame,style="Card.TFrame",padding=10); card.pack(fill=X,pady=(0,8)); ttk.Label(card,text=text,style="Section.TLabel").pack(anchor="w"); ttk.Label(card,text=desc,style="Muted.TLabel",wraplength=620).pack(anchor="w"); ttk.Button(card,text=_tr('ui.source.start.952f3754'),command=lambda k=kind:run(k)).pack(anchor="e")
            ttk.Label(frame,textvariable=result,wraplength=650).pack(anchor="w",pady=(8,0)); ttk.Button(frame,text=_tr('ui.source.open.back.upmap.acedffd2'),command=self._open_backup_dir).pack(anchor="e",pady=(8,0))
        def _open_project_roadmap(self) -> None:
            projects = self._get_selected_projects()
            project = projects[0] if projects else (self.selected_project if self.selected_project else None)
            if not project:
                return
            path = project.path / "ROADMAP.md"
            if not path.exists():
                messagebox.showinfo(_tr('ui.source.roadmap.8119e7f3'), _tr('ui.source.roadmap.md.ontbreekt.start.eerst.project.archi.84432253'))
                return
            windows_open_path(path)
        def _v92_save_version_snapshot(self) -> None:
            try:
                payload=self._v90_report_payload()
                case_dir=self._v90_case_dir(payload.get("case_id","CASE"))
                versions=case_dir/"reports"/"versions"; versions.mkdir(parents=True,exist_ok=True)
                stamp=_dt.datetime.now().strftime("%Y%m%d_%H%M%S")
                title=re.sub(r"[^A-Za-z0-9_.-]+","_",payload.get("title","rapport"))[:60]
                path=versions/f"{title}_{stamp}.pmreport.json"
                path.write_text(json.dumps(payload,indent=2,ensure_ascii=False),encoding="utf-8")
                sha=hashlib.sha256(path.read_bytes()).hexdigest()
                path.with_suffix(path.suffix+".sha256").write_text(f"{sha}  {path.name}\\n",encoding="utf-8")
                messagebox.showinfo(_tr('ui.source.versiesnapshot.a7cbee96'),_tr('ui.source.versiesnapshot.opgeslagen.n.p0.58a84d80',p0=path),parent=getattr(self,"_report_studio",self.root))
            except Exception as exc:
                messagebox.showerror(_tr('ui.source.versiesnapshot.mislukt.7ef388a8'),str(exc),parent=getattr(self,"_report_studio",self.root))
        def _v93_export_docx(self) -> None:
            try:
                from docx import Document
                from docx.shared import Inches
            except Exception:
                messagebox.showwarning(_tr('ui.source.docx.export.0ddc2686'),_tr('ui.source.python.docx.is.niet.ge.nstalleerd.installeer.m.b86f8e7f'),parent=getattr(self,"_report_studio",self.root)); return
            payload=self._v90_report_payload(); case_dir=self._v90_case_dir(payload["case_id"])/"exports"; case_dir.mkdir(parents=True,exist_ok=True)
            safe=re.sub(r"[^A-Za-z0-9_.-]+","_",payload["title"])[:80] or "rapport"
            chosen=filedialog.asksaveasfilename(title=_tr('ui.source.export.docx.73c8f571'),initialdir=str(case_dir),initialfile=safe+".docx",defaultextension=".docx",filetypes=[(_tr('ui.source.word.document.70a80dfa'),"*.docx")])
            if not chosen:return
            doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(0.75); sec.bottom_margin=Inches(0.75)
            doc.add_heading(payload["title"],0); doc.add_paragraph(f"Case: {payload['case_id']} | Classificatie: {payload['classification']}")
            for line in payload["content"].splitlines():
                if line.startswith("### "): doc.add_heading(line[4:],level=3)
                elif line.startswith("## "): doc.add_heading(line[3:],level=2)
                elif line.startswith("# "): doc.add_heading(line[2:],level=1)
                elif line.startswith("- "): doc.add_paragraph(line[2:],style="List Bullet")
                elif line.strip()=="--- PAGINAEINDE ---": doc.add_page_break()
                else: doc.add_paragraph(line)
            path=Path(chosen); doc.save(path); sha=hashlib.sha256(path.read_bytes()).hexdigest(); path.with_suffix(path.suffix+".sha256").write_text(f"{sha}  {path.name}\n",encoding="utf-8")
            self._v93_add_audit("DOCX-export",str(path)); messagebox.showinfo(_tr('ui.source.docx.export.0ddc2686'),_tr('ui.source.docx.opgeslagen.p0.8ae4e2b9',p0=path),parent=getattr(self,"_report_studio",self.root))
